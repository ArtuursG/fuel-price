"""Parser for eu-weekly-oil-bulletin (docs/sources.yaml). Pure function:
XLSX bytes in, OfficialWeeklyPrice records out.

File structure (confirmed 2026-09-19): single sheet, row 1 = product column
headers (French/English/German), row 2 = bulletin reference date (col A) and
units (EUR per "1000 l" for the road-fuel columns), rows 3+ = one row per
country. Prices are EUR per 1000 litres, WITH taxes -- numerically identical
to our price_milli convention (milli-EUR per litre), just needs rounding to
an int. Only LV/LT/EE rows are kept (see docs/DECISIONS.md: cross-border
scope is Baltic, not full EU) and only road-fuel columns we have product
codes for (P95, DSL, LPG) -- heating oil/fuel oil columns are dropped.

The "merchant" field on OfficialWeeklyPrice (see core/models.py) holds a
country code here, not a company name -- the column is a generic label, and
this is the only official_weekly-shaped source implemented so far.
"""

from __future__ import annotations

import io
import logging

from openpyxl import load_workbook

from collector.core.models import OfficialWeeklyPrice

logger = logging.getLogger(__name__)

SOURCE_ID = "eu-weekly-oil-bulletin"
URL = (
    "https://energy.ec.europa.eu/document/download/"
    "264c2d0f-f161-4ea3-a777-78faae59bea0_en"
    "?filename=Weekly%20Oil%20Bulletin%20Weekly%20prices%20with%20Taxes%20-%202024-02-19.xlsx"
)

COUNTRY_MAP = {
    "Latvia": "LV",
    "Lithuania": "LT",
    "Estonia": "EE",
}

# Substrings matched against row-1 header text (case-sensitive, as published).
PRODUCT_HEADER_MARKERS = {
    "P95": "Euro-super 95",
    "DSL": "Gas oil automobile",
    "LPG": "GPL pour moteur",
}


def _find_product_columns(header_row: tuple) -> dict[str, int]:
    columns: dict[str, int] = {}
    for col_idx, cell in enumerate(header_row):
        if not cell:
            continue
        for product, marker in PRODUCT_HEADER_MARKERS.items():
            if marker in str(cell):
                columns[product] = col_idx
    return columns


def parse(content: bytes) -> list[OfficialWeeklyPrice]:
    wb = load_workbook(io.BytesIO(content), data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 3:
        return []

    header_row, date_row = rows[0], rows[1]
    product_columns = _find_product_columns(header_row)
    missing = set(PRODUCT_HEADER_MARKERS) - set(product_columns)
    if missing:
        logger.warning("EU Weekly Oil Bulletin: could not find columns for %s", missing)

    bulletin_date = date_row[0]
    if bulletin_date is None:
        logger.warning("EU Weekly Oil Bulletin: no bulletin date in row 2")
        return []
    week_monday = (
        bulletin_date.date().isoformat()
        if hasattr(bulletin_date, "date")
        else str(bulletin_date)[:10]
    )

    prices: list[OfficialWeeklyPrice] = []
    for row in rows[2:]:
        if not row or not row[0]:
            continue
        country = COUNTRY_MAP.get(str(row[0]).strip())
        if country is None:
            continue  # not LV/LT/EE, or a footer/average row

        for product, col_idx in product_columns.items():
            if col_idx >= len(row) or row[col_idx] is None:
                continue
            prices.append(
                OfficialWeeklyPrice(
                    week_monday=week_monday,
                    merchant=country,
                    product=product,
                    avg_price_milli=round(float(row[col_idx])),
                    source_url=URL,
                )
            )

    return prices
